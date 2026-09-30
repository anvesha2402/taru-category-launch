"""D7 Claims substantiation matrix. Every public TARU claim, its evidence and governing clause.
Outputs compliance/TARU_Claims_Matrix.xlsx and compliance/claims_matrix.csv (input to notebook 07, section 10)."""
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
C = [
# id, wording, where, type, evidence needed, evidence status, clause, risk, decision
("CL01","Organic cotton fabric, certified by [certification body], licence no. [X]. Scan to see the certificate and what it covers.","Swing tag; product page","Certification","Valid scope + transaction certificates for the fabric and every processing stage","Not held (fictional brand)","CCPA 1 [part 2]; CCPA 6(1); ASCI-GREEN 6; GOTS label grade ≥95%","Medium","Conditional: use only once certificates exist"),
("CL02","Made with [X]% organic cotton, licence no. [X].","Swing tag","Certification","Transaction certificate stating the organic %; GOTS 'made with' grade ≥70%","Not held","CCPA 1 [part 2]; CCPA 6(1); GOTS label grades","Medium","Conditional"),
("CL03","Scan to see where your fabric was grown and woven.","Swing tag; QR page","Traceability","Named farm, gin, mill, garment units, each with a certificate reference","Not held","CCPA 6(1)","Low","Conditional"),
("CL04","The fabric is certified organic; the packaging is not covered by this certification.","Product page","Scope statement","Certificate scope","Not held","CCPA 6(3); ASCI-GREEN 4 [part 2]","Low","Conditional"),
("CL05","Fibre: 100% cotton. Care: gentle wash, dry in shade. Made in India.","Care label","Mandatory label","Fibre test report; care test","Available once sampled","TC-DRAFT 7; TC-DRAFT 1 [part 3]; Legal Metrology","Low","Use"),
("CL06","Cut with 4 cm more room through the waist than our standard block.","Product page; fit guide","Fit","Documented pattern grading","Achievable in sampling","No environmental claim","Low","Use"),
("CL07","Free alterations within 30 days of delivery.","Product page; box insert","Service","Published service terms","Policy decision","No environmental claim","Low","Use"),
("CL08","Soft through a six-hour function.","Campaign; survey Card C","Performance","Documented wear trial with method and sample size","Not held","CCPA 1 [part 2] (verifiable evidence); ASCI Code Ch. I (substantiation)","High","Test-only (survey stimulus); do not publish without trial"),
("CL09","Rooted in nature. Refined by time.","Tagline candidate","Brand line","Qualifier or substantiation for 'nature'","None","CCPA 5 (generic terms); CCPA 2 [part 2]","Medium","On hold"),
("CL10","Know what you wear.","Tagline candidate","Brand line","None: no environmental or performance claim","n/a","No environmental claim","Low","Use"),
("CL11","Every thread, accounted for.","Tagline candidate","Traceability","Trail covering all fibres including sewing thread","Not held","CCPA 6(1); CCPA 4 [part 2]","Medium","Conditional"),
("CL12","Made from natural fibres.","Survey Card A (control)","Generic","n/a (control stimulus)","n/a","CCPA 5 ('natural' is generic)","High","Test-only; never publish"),
("CL13","GOTS-certified organic cotton. Licence no. XX-GOTS-000000","Survey Card B","Certification","n/a (placeholder licence)","n/a","CCPA 6 (placeholder must never look real in public)","High","Test-only; never publish"),
("CL14","Eco-friendly / sustainable / green / planet-friendly (unqualified)","Anywhere","Generic","Not substantiable as an absolute claim","n/a","CCPA 5; ASCI-GREEN 1 [part 2]","High","Banned"),
("CL15","100% natural / chemical-free / non-toxic / hypoallergenic","Anywhere","Absolute / health","Not substantiable without full testing","n/a","CCPA 6(5); ASCI-GREEN 10","High","Banned"),
("CL16","India's first / only certified-organic menswear brand","Anywhere","Comparative","Untrue: Isha Life sells an organic cotton men's kurta (register R032)","n/a","CCPA 6(4); ASCI-GREEN 2 [part 2]","High","Banned"),
("CL17","Organic bamboo / organic lyocell","Anywhere","Fibre","Processed fibres cannot be organic","n/a","GOTS label grades; CCPA 5","High","Banned"),
("CL18","Lookbook image created using AI.","Social; website","AI disclosure","Label on every AI image","Policy","ASCI-AI 2; ASCI-AI 4; IT Amendment Rules 2026","Low","Use (mandatory on AI images)"),
("CL19","AI-generated customer testimonial or review","Anywhere","Endorsement","Prohibited even with a label","n/a","ASCI-AI 1","High","Banned"),
("CL20","This box is made from recycled paperboard, certified by [body]; applies to the box only.","Packaging","Packaging","FSC or recycled-content certificate","Not held","ASCI-GREEN 4 [part 2]; ASCI-GREEN 6; CCPA 6(3)","Low","Conditional"),
("CL21","Our goal: [X]% certified-organic fibre across the range by [year]; plan and progress published at [link].","About page","Aspirational","Published actionable plan with milestones","Not held","ASCI-GREEN 8","Medium","Conditional"),
("CL22","No markdowns on the Ceremonial line.","Store; website","Pricing","Pricing policy","Policy decision","Legal Metrology (MRP display)","Low","Use"),
]
cols=["claim_id","claim_wording","where_used","claim_type","evidence_required","evidence_status","governing_clause","risk","decision","guardrail_verdict","guardrail_clause","reviewed_by","review_date"]
df=pd.DataFrame([list(c)+["","","",""] for c in C],columns=cols)
df.to_csv("compliance/claims_matrix.csv",index=False)
P="compliance/TARU_Claims_Matrix.xlsx"; df.to_excel(P,index=False,sheet_name="Claims_Matrix")
wb=load_workbook(P); ws=wb.active
HB=Font(name="Arial",bold=True,color="FFFFFF"); HF=PatternFill("solid",start_color="2F4F3E")
fills={"Use":"D9EAD3","Conditional":"FFF2CC","Test-only":"CFE2F3","On hold":"FCE5CD","Banned":"F4CCCC"}
for c in ws[1]: c.font=HB; c.fill=HF; c.alignment=Alignment(wrap_text=True,vertical="center")
for w,col in zip([8,48,20,14,34,18,34,8,30,16,18,12,11],"ABCDEFGHIJKLM"): ws.column_dimensions[col].width=w
for row in ws.iter_rows(min_row=2):
    for c in row: c.font=Font(name="Arial",size=10); c.alignment=Alignment(wrap_text=True,vertical="top")
    d=row[8].value
    for k,v in fills.items():
        if d.startswith(k): row[8].fill=PatternFill("solid",start_color=v)
dv=DataValidation(type="list",formula1='"Use,Conditional,Test-only,On hold,Banned"',allow_blank=False); ws.add_data_validation(dv); dv.add("I2:I200")
ws.freeze_panes="C2"
r=ws.max_row+2
for i,t in enumerate(["Rule: no claim appears in any TARU material unless it has a row here marked Use or Conditional (with evidence in hand).",
 "Guardrail columns are filled by notebooks/07_claims_guardrail.ipynb, section 10. The guardrail recommends; the author decides.",
 "TARU is a fictional brand for an independent student portfolio project; 'Not held' means the evidence would be required in a real launch."]):
    ws.cell(r+i,1,t).font=Font(name="Arial",size=9,italic=True)
wb.save(P); print(df.decision.str.split(":").str[0].str.split(" ").str[0].value_counts().to_dict())
