"""The app's model must reproduce the spreadsheet (model/taru_economics.xlsx, cached values)
and the pivot result (model/sensitivity_results.json)."""
import json, sys
from pathlib import Path
import pytest
ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "app"))
import taru_model as m  # noqa: E402

def xlsx():
    from openpyxl import load_workbook
    return load_workbook(ROOT / "model" / "taru_economics.xlsx", data_only=True)

def test_cogs_and_net_price_match_range_cost():
    R = xlsx()["Range_Cost"]; p = m.params()
    for i, line in enumerate(m.LINES):
        assert m.net_price(line, p) == pytest.approx(R.cell(5 + i, 4).value)
        assert m.cogs(line, p) == pytest.approx(R.cell(5 + i, 11).value)

def test_every_waterfall_cell_matches():
    W = xlsx()["Channel_Waterfalls"]; p = m.params()
    col = 2
    for line in m.LINES:
        for ch in m.CHANNELS[:4]:
            assert m.contribution(line, ch, p) == pytest.approx(W.cell(17, col).value), (line, ch)
            assert m.revenue(line, ch, p) == pytest.approx(W.cell(7, col).value), (line, ch)
            col += 1

def test_blended_and_breakeven_match_plan_as_briefed():
    W = xlsx()["Channel_Waterfalls"]; p = m.params()
    b = m.blended(p)
    assert b["contribution"] == pytest.approx(W["B25"].value)
    assert b["revenue"] == pytest.approx(W["B22"].value)
    res = json.loads((ROOT / "model" / "sensitivity_results.json").read_text())
    assert m.breakeven_sets_per_month(p) == pytest.approx(res["base"]["breakeven_sets_m"])

def test_pivot_matches_sensitivity_results():
    res = json.loads((ROOT / "model" / "sensitivity_results.json").read_text())["pivot"]
    p = m.params("pivot")
    assert m.blended(p)["contribution"] == pytest.approx(res["cm_per_set"])
    assert m.breakeven_sets_per_month(p) == pytest.approx(res["breakeven_sets_m"])

def test_gst_slab_switches_above_2500():
    p = m.params(); p["line"]["Everyday"]["price"] = 2600
    assert m.gst_rate(2600, p) == 0.18 and m.net_price("Everyday", p) == pytest.approx(2600 / 1.18)

def test_retail_partner_off_by_default_and_costs_when_on():
    p = m.params(); assert m.channel_mix(p)["Retail partner"] == 0
    steps = dict(m.waterfall("Festive", "Retail partner", p))
    assert steps["Retailer margin"] < 0 and steps["Sale-or-return stock back"] < 0

def test_breakeven_never_when_contribution_negative():
    p = m.params()
    for c in m.CHANNELS: p["channel"][c]["mix"] = 0
    p["channel"]["Marketplace"]["mix"] = 1
    p["line"]["Everyday"]["mix"], p["line"]["Festive"]["mix"], p["line"]["Ceremonial"]["mix"] = 1, 0, 0
    assert m.breakeven_sets_per_month(p) is None

def test_vw_warning():
    p = m.params(); assert m.vw_check("Everyday", p)[0] == "ok"
    p["line"]["Everyday"]["price"] = 3999; assert m.vw_check("Everyday", p)[0] == "outside"
    assert m.vw_check("Festive", p)[0] == "none"
